#!/usr/bin/env python3
"""
Renders the report to DOCX (OOXML).

    python3 render_docx.py

Output: Deccan_OT-03_Incident_and_Security_Report.docx
"""

from __future__ import annotations

import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

from content_p1 import BLOCKS_P1
from content_p2 import BLOCKS_P2
from content_p3 import BLOCKS_P3
from content_p4 import BLOCKS_P4

BLOCKS = BLOCKS_P1 + BLOCKS_P2 + BLOCKS_P3 + BLOCKS_P4

INK = RGBColor(0x1B, 0x1F, 0x24)
DIM = RGBColor(0x5A, 0x63, 0x6E)
BAD = RGBColor(0xA3, 0x2C, 0x22)
ACC = RGBColor(0x1F, 0x4E, 0x79)
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"


def shade(cell_or_par, hex_fill: str) -> None:
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def set_cell_width(cell, cm: float) -> None:
    cell.width = Cm(cm)


def para(doc, text="", size=10.5, bold=False, italic=False, color=INK,
         space_after=6, space_before=0, align=None, font=BODY_FONT, indent=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if align:
        p.alignment = align
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    return p


def add_bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(it)
        run.font.size = Pt(10.5)
        run.font.name = BODY_FONT
        run.font.color.rgb = INK
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(0.7)


def add_numbers(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(it)
        run.font.size = Pt(10.5)
        run.font.name = BODY_FONT
        run.font.color.rgb = INK
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(0.7)


def add_table(doc, spec):
    rows = spec["rows"]
    ncol = len(spec["cols"]) or len(rows[0])
    header = spec.get("cols") or [""] * ncol
    show_header = spec.get("cols") and any(c.strip() for c in spec["cols"])
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = spec.get("widths") or [1.0 / ncol] * ncol
    total_cm = 17.0

    def fill_row(values, bold=False, fill=None, size=9.0):
        cells = t.add_row().cells
        for i, val in enumerate(values):
            if i >= ncol:
                break
            set_cell_width(cells[i], total_cm * widths[i])
            cp = cells[i].paragraphs[0]
            cp.paragraph_format.space_after = Pt(2)
            cp.paragraph_format.space_before = Pt(2)
            run = cp.add_run(str(val))
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = BODY_FONT
            run.font.color.rgb = INK
            if fill:
                shade(cells[i], fill)

    if show_header:
        fill_row(header, bold=True, fill="EDF1F5", size=9.0)
    for r in rows:
        fill_row(r, size=8.8 if ncol > 3 else 9.2)
    if spec.get("caption"):
        para(doc, spec["caption"], size=8.6, italic=True, color=DIM, space_before=4)
    return t


def add_kv(doc, pairs):
    t = doc.add_table(rows=0, cols=2)
    for k, v in pairs:
        cells = t.add_row().cells
        set_cell_width(cells[0], 3.4)
        set_cell_width(cells[1], 13.6)
        r1 = cells[0].paragraphs[0].add_run(k)
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.name = BODY_FONT
        r2 = cells[1].paragraphs[0].add_run(v)
        r2.font.size = Pt(9.5)
        r2.font.name = BODY_FONT
        for c in cells:
            c.paragraphs[0].paragraph_format.space_after = Pt(3)
    return t


def add_code(doc, text):
    for line in text.split("\n"):
        p = doc.add_paragraph()
        run = p.add_run(line if line else " ")
        run.font.size = Pt(8.0)
        run.font.name = MONO_FONT
        run.font.color.rgb = RGBColor(0x22, 0x2A, 0x33)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Cm(0.25)
        p.paragraph_format.line_spacing = 1.0
        shade(p, "F4F6F8")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_callout(doc, text):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.rows[0].cells[0]
    set_cell_width(cell, 17.0)
    shade(cell, "FBF3E4")
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.name = BODY_FONT
    run.font.color.rgb = RGBColor(0x6B, 0x4E, 0x0B)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc, spec):
    path = os.path.join(ROOT, "diagrams", spec["path"])
    if os.path.exists(path):
        doc.add_picture(path, width=Cm(16.6))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, spec["caption"], size=8.6, italic=True, color=DIM,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)


def add_heading(doc, text, level):
    sizes = {1: 15.5, 2: 12.2, 3: 11.0}
    p = doc.add_paragraph()
    if level == 1:
        p.paragraph_format.space_before = Pt(10)
    else:
        p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.font.size = Pt(sizes[level])
    run.font.bold = True
    run.font.name = BODY_FONT
    run.font.color.rgb = ACC if level == 1 else INK
    if level == 1:
        # keep each top level section starting on its own page
        p.paragraph_format.page_break_before = True
    return p


def build() -> str:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    for attr, val in (("top_margin", 1.9), ("bottom_margin", 1.9),
                      ("left_margin", 2.0), ("right_margin", 2.0)):
        setattr(sec, attr, Cm(val))

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(10.5)

    # footer with page numbers
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Deccan OT-03  -  incident and security analysis  -  page ")
    fr.font.size = Pt(8)
    fr.font.color.rgb = DIM
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    first_h1 = True
    prev = None
    for block in BLOCKS:
        kind = block[0]
        if kind == "pagebreak" and prev == "pagebreak":
            continue
        prev = kind
        if kind == "cover":
            spec = block[1]
            para(doc, spec["kicker"].upper(), size=9.5, bold=True, color=DIM,
                 space_before=40, space_after=8)
            para(doc, spec["title"], size=24, bold=True, color=ACC, space_after=2)
            para(doc, spec["subtitle"], size=13, color=INK, space_after=18)
            t = doc.add_table(rows=0, cols=2)
            for k, v in spec["meta"]:
                cells = t.add_row().cells
                set_cell_width(cells[0], 3.6)
                set_cell_width(cells[1], 13.4)
                r1 = cells[0].paragraphs[0].add_run(k.upper())
                r1.font.size = Pt(8.2)
                r1.font.bold = True
                r1.font.color.rgb = DIM
                r2 = cells[1].paragraphs[0].add_run(v)
                r2.font.size = Pt(9.6)
                for c in cells:
                    c.paragraphs[0].paragraph_format.space_after = Pt(3)
            doc.add_paragraph()
            add_callout(doc, spec["note"])
            para(doc, "", space_after=0)
        elif kind == "h1":
            p = add_heading(doc, block[1], 1)
            if first_h1:
                p.paragraph_format.page_break_before = False
                first_h1 = False
        elif kind == "h2":
            add_heading(doc, block[1], 2)
        elif kind == "h3":
            add_heading(doc, block[1], 3)
        elif kind == "p":
            para(doc, block[1], align=WD_ALIGN_PARAGRAPH.LEFT, space_after=7)
        elif kind == "bullets":
            add_bullets(doc, block[1])
        elif kind == "numbered":
            add_numbers(doc, block[1])
        elif kind == "table":
            add_table(doc, block[1])
        elif kind == "kv":
            add_kv(doc, block[1])
        elif kind == "code":
            add_code(doc, block[1])
        elif kind == "callout":
            add_callout(doc, block[1])
        elif kind == "figure":
            add_figure(doc, block[1])
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        else:
            raise ValueError(f"unknown block: {kind}")

    out = os.path.join(ROOT, "Deccan_OT-03_Incident_and_Security_Report.docx")
    doc.save(out)
    print("wrote", out)
    return out


if __name__ == "__main__":
    build()
