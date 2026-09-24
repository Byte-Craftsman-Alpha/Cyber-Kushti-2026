# -*- coding: utf-8 -*-
"""
build_report.py -- renders report_content.py to DOCX and PDF.

    python3 build_report.py

Both outputs come from the same block list, so the two documents cannot drift.
"""

import os
import sys
from xml.sax.saxutils import escape

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import report_content as RC

OUT_DOCX = os.path.join(HERE, "SUP-01-Anvil-Incident-Report.docx")
OUT_PDF = os.path.join(HERE, "SUP-01-Anvil-Incident-Report.pdf")

A4_W, A4_H = 595.27, 841.89
MARGIN = 56
CONTENT_W = A4_W - 2 * MARGIN

DARK = (0.23, 0.25, 0.27)
GREY = (0.42, 0.45, 0.48)


def transcript_lines(path):
    full = os.path.join(HERE, path) if not os.path.isabs(path) else path
    if not os.path.exists(full):
        return ["(transcript not found -- run prototype/run_simulation.py to generate it)"]
    out = []
    for line in open(full, encoding="utf-8", errors="replace").read().splitlines():
        line = line.rstrip()
        if len(line) <= 104:
            out.append(line)
        else:
            while len(line) > 104:
                out.append(line[:104])
                line = "    " + line[104:]
            out.append(line)
    return out


def numbered_blocks():
    """Yield (kind, kc_number, args, kw) so kill-chain steps can be numbered."""
    n = 0
    for kind, args, kw in RC.BLOCKS:
        if kind == "kc":
            n += 1
            yield kind, n, args, kw
        else:
            yield kind, 0, args, kw


# ===========================================================================
# DOCX
# ===========================================================================
def build_docx():
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin = sec.right_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.13

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), hexcolor)
        tcPr.append(shd)

    def para(text, size=10.5, bold=False, italic=False, color=None, align=None,
             font="Georgia", space_after=7, space_before=0, indent=0):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.name = font
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        if color:
            r.font.color.rgb = RGBColor(*[int(c * 255) for c in color])
        if align:
            p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        if indent:
            p.paragraph_format.left_indent = Cm(indent)
        return p

    def mono_block(text, size=8):
        for line in text.splitlines():
            p = doc.add_paragraph()
            r = p.add_run(line if line else " ")
            r.font.name = "Consolas"
            r.font.size = Pt(size)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0

    for kind, n, args, kw in numbered_blocks():
        if kind == "title":
            para(args[0], size=27, bold=True, space_before=36, space_after=4)
        elif kind == "subtitle":
            para(args[0], size=15, italic=True, color=GREY, space_after=18)
        elif kind == "deck":
            para(args[0], size=12, space_after=14)
        elif kind == "meta":
            para(args[0], size=10, italic=True, color=GREY, space_after=6)
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "h1":
            para(args[0], size=16, bold=True, space_before=16, space_after=8, color=DARK)
        elif kind == "h2":
            para(args[0], size=12.5, bold=True, space_before=12, space_after=5, color=DARK)
        elif kind == "p":
            para(args[0])
        elif kind == "caption":
            para(args[0], size=8.5, italic=True, color=GREY,
                 align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
        elif kind == "bullets":
            for b in args[0]:
                p = doc.add_paragraph(style="List Bullet")
                r = p.add_run(b)
                r.font.name = "Georgia"
                r.font.size = Pt(10.5)
                p.paragraph_format.space_after = Pt(3)
        elif kind == "numbers":
            for i, b in enumerate(args[0], 1):
                p = doc.add_paragraph()
                r = p.add_run(f"{i}.  ")
                r.font.name = "Georgia"; r.font.size = Pt(10.5); r.bold = True
                r2 = p.add_run(b)
                r2.font.name = "Georgia"; r2.font.size = Pt(10.5)
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_after = Pt(4)
        elif kind == "legend":
            for tag, name, desc in args[0]:
                p = doc.add_paragraph()
                r = p.add_run(f"{tag}  {name} — ")
                r.font.name = "Georgia"; r.font.size = Pt(10.5); r.bold = True
                r2 = p.add_run(desc)
                r2.font.name = "Georgia"; r2.font.size = Pt(10.5)
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_after = Pt(3)
        elif kind == "table":
            cols, widths, rows = kw["cols"], kw["widths"], kw["rows"]
            t = doc.add_table(rows=1 + len(rows), cols=len(cols))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            total = sum(widths)
            usable = 17.0
            for j, c in enumerate(cols):
                cell = t.rows[0].cells[j]
                cell.text = ""
                p = cell.paragraphs[0]
                r = p.add_run(c)
                r.bold = True
                r.font.size = Pt(9)
                r.font.name = "Georgia"
                shade(cell, "E8E5DE")
                cell.width = Cm(usable * widths[j] / total)
            for i, row in enumerate(rows):
                for j, val in enumerate(row):
                    cell = t.rows[1 + i].cells[j]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    r = p.add_run(val)
                    r.font.size = Pt(8.5)
                    r.font.name = "Georgia"
                    p.paragraph_format.space_after = Pt(1)
                    cell.width = Cm(usable * widths[j] / total)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif kind == "code":
            mono_block(args[0])
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif kind == "image":
            doc.add_picture(os.path.join(HERE, args[0]), width=Cm(17))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind == "kc":
            title, command, details = args
            para(f"Step {n} — {title}", size=12, bold=True, space_before=12,
                 space_after=4, color=DARK)
            p = doc.add_paragraph()
            r = p.add_run("Command or tool: ")
            r.font.name = "Georgia"; r.font.size = Pt(9.5); r.bold = True
            if command.strip():
                r2 = p.add_run(command)
                r2.font.name = "Consolas"; r2.font.size = Pt(8.5)
            else:
                r2 = p.add_run("(none — the victim's own tooling does the work)")
                r2.font.name = "Georgia"; r2.font.size = Pt(9.5); r2.italic = True
            p.paragraph_format.space_after = Pt(3)
            p2 = doc.add_paragraph()
            r = p2.add_run("Details: ")
            r.font.name = "Georgia"; r.font.size = Pt(10.5); r.bold = True
            r2 = p2.add_run(details)
            r2.font.name = "Georgia"; r2.font.size = Pt(10.5)
            p2.paragraph_format.space_after = Pt(6)
        elif kind == "transcript":
            mono_block("\n".join(transcript_lines(args[0])), size=7)
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.save(OUT_DOCX)
    print("wrote", OUT_DOCX)


# ===========================================================================
# PDF
# ===========================================================================
def build_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, Preformatted, Image, PageBreak,
                                    KeepTogether)
    from reportlab.pdfbase import pdfmetrics

    styles = {
        "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=25,
                                leading=29, textColor=colors.HexColor("#3a3f45"),
                                spaceBefore=40, spaceAfter=6),
        "subtitle": ParagraphStyle("subtitle", fontName="Times-Italic", fontSize=14,
                                   leading=18, textColor=colors.HexColor("#6b7076"),
                                   spaceAfter=16),
        "deck": ParagraphStyle("deck", fontName="Times-Roman", fontSize=11.5,
                               leading=15.5, spaceAfter=12),
        "meta": ParagraphStyle("meta", fontName="Times-Italic", fontSize=9.5,
                               leading=12, textColor=colors.HexColor("#6b7076"),
                               spaceAfter=4),
        "h1": ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=14.5,
                             leading=18, textColor=colors.HexColor("#3a3f45"),
                             spaceBefore=16, spaceAfter=7),
        "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=11.5,
                             leading=15, textColor=colors.HexColor("#3a3f45"),
                             spaceBefore=11, spaceAfter=4),
        "body": ParagraphStyle("body", fontName="Times-Roman", fontSize=10,
                               leading=13.8, spaceAfter=7),
        "caption": ParagraphStyle("caption", fontName="Times-Italic", fontSize=8.5,
                                  leading=11, alignment=1, textColor=colors.HexColor("#6b7076"),
                                  spaceAfter=9),
        "bullet": ParagraphStyle("bullet", fontName="Times-Roman", fontSize=10,
                                 leading=13.5, leftIndent=14, bulletIndent=2,
                                 spaceAfter=3),
        "num": ParagraphStyle("num", fontName="Times-Roman", fontSize=10,
                              leading=13.5, leftIndent=18, spaceAfter=4),
        "legend": ParagraphStyle("legend", fontName="Times-Roman", fontSize=10,
                                 leading=13.5, leftIndent=14, spaceAfter=3),
        "kctitle": ParagraphStyle("kctitle", fontName="Helvetica-Bold", fontSize=11.5,
                                  leading=15, textColor=colors.HexColor("#3a3f45"),
                                  spaceBefore=11, spaceAfter=3),
        "kcmeta": ParagraphStyle("kcmeta", fontName="Times-Roman", fontSize=9.5,
                                 leading=12.5, spaceAfter=3, wordWrap="CJK"),
        "kcdet": ParagraphStyle("kcdet", fontName="Times-Roman", fontSize=10,
                                leading=13.8, spaceAfter=6),
        "code": ParagraphStyle("code", fontName="Courier", fontSize=7.8,
                               leading=9.8, spaceAfter=6,
                               backColor=colors.HexColor("#f4f2ed")),
        "transcript": ParagraphStyle("transcript", fontName="Courier", fontSize=6.5,
                                     leading=8.1, spaceAfter=0),
        "tbl": ParagraphStyle("tbl", fontName="Times-Roman", fontSize=7.8,
                              leading=9.6),
        "tblh": ParagraphStyle("tblh", fontName="Helvetica-Bold", fontSize=8.2,
                               leading=10),
    }

    def footer(canvas, docobj):
        canvas.saveState()
        canvas.setFont("Times-Italic", 8)
        canvas.setFillColor(colors.HexColor("#8a9199"))
        canvas.drawString(MARGIN, 28, "SUP-01 Anvil Systems — incident analysis")
        canvas.drawRightString(A4_W - MARGIN, 28, f"page {docobj.page}")
        canvas.restoreState()

    pdf = SimpleDocTemplate(OUT_PDF, pagesize=A4,
                            leftMargin=MARGIN, rightMargin=MARGIN,
                            topMargin=MARGIN, bottomMargin=MARGIN,
                            title="SUP-01 Anvil Systems — incident analysis",
                            author="Incident analysis pack")

    story = []
    for kind, n, args, kw in numbered_blocks():
        if kind == "title":
            story.append(Paragraph(escape(args[0]), styles["title"]))
        elif kind == "subtitle":
            story.append(Paragraph(escape(args[0]), styles["subtitle"]))
        elif kind == "deck":
            story.append(Paragraph(escape(args[0]), styles["deck"]))
        elif kind == "meta":
            story.append(Paragraph(escape(args[0]), styles["meta"]))
            story.append(PageBreak())
        elif kind == "h1":
            story.append(Paragraph(escape(args[0]), styles["h1"]))
        elif kind == "h2":
            story.append(Paragraph(escape(args[0]), styles["h2"]))
        elif kind == "p":
            story.append(Paragraph(escape(args[0]), styles["body"]))
        elif kind == "caption":
            story.append(Paragraph(escape(args[0]), styles["caption"]))
        elif kind == "bullets":
            for b in args[0]:
                story.append(Paragraph(escape(b), styles["bullet"], bulletText="•"))
        elif kind == "numbers":
            for i, b in enumerate(args[0], 1):
                story.append(Paragraph(escape(b), styles["num"], bulletText=f"{i}."))
        elif kind == "legend":
            for tag, name, desc in args[0]:
                story.append(Paragraph(
                    f'<b>{escape(tag)}  {escape(name)} —</b> {escape(desc)}',
                    styles["legend"]))
        elif kind == "table":
            cols, widths, rows = kw["cols"], kw["widths"], kw["rows"]
            total = sum(widths)
            col_w = [CONTENT_W * w / total for w in widths]
            data = [[Paragraph(escape(c), styles["tblh"]) for c in cols]]
            for row in rows:
                data.append([Paragraph(escape(str(v)), styles["tbl"]) for v in row])
            t = Table(data, colWidths=col_w, repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e5de")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#b8bcc0")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(t)
            story.append(Spacer(1, 8))
        elif kind == "code":
            story.append(Preformatted(args[0], styles["code"]))
        elif kind == "image":
            from PIL import Image as PILImage
            img_path = os.path.join(HERE, args[0])
            iw, ih = PILImage.open(img_path).size
            w = CONTENT_W
            h = w * ih / iw
            story.append(Image(img_path, width=w, height=h))
        elif kind == "kc":
            title, command, details = args
            story.append(Paragraph(f"Step {n} — {escape(title)}", styles["kctitle"]))
            if command.strip():
                story.append(Paragraph(
                    f'<b>Command or tool:</b> <font face="Courier" size="8">{escape(command)}</font>',
                    styles["kcmeta"]))
            else:
                story.append(Paragraph(
                    "<b>Command or tool:</b> <i>(none — the victim's own tooling does the work)</i>",
                    styles["kcmeta"]))
            story.append(Paragraph(f"<b>Details:</b> {escape(details)}", styles["kcdet"]))
        elif kind == "transcript":
            story.append(Preformatted("\n".join(transcript_lines(args[0])), styles["transcript"]))
        elif kind == "pagebreak":
            story.append(PageBreak())

    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    print("wrote", OUT_PDF)


if __name__ == "__main__":
    build_docx()
    build_pdf()
